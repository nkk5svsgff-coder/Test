#!/usr/bin/env python3
"""Write Software_Meta_Advertisers_FINAL_with_Stripe.xlsx

Inserts a "Stripe?" column immediately to the right of Website, and appends
audit columns (evidence / other payment tech / check status / URL checked) at
the far right. Rebuilds the sheet cell-by-cell because openpyxl's insert_cols
corrupts hyperlinks (this sheet has 30k+ of them).
"""

import json
import os
from copy import copy

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SRC = "/root/.claude/uploads/d588e071-b9c0-5f19-99b9-191f0e20e61d/b4137442-Software_Meta_Advertisers_FINAL.xlsx"
HERE = os.path.dirname(os.path.abspath(__file__))
D = os.path.join(HERE, "data")
OUT = os.path.join(D, "Software_Meta_Advertisers_FINAL_with_Stripe.xlsx")

SHEET = "Software Advertisers"
WEBSITE_COL = 3          # source column C
INSERT_AT = 4            # new "Stripe?" column becomes D


def norm(u):
    u = (u or "").strip()
    if not u:
        return ""
    if not u.startswith(("http://", "https://")):
        u = "https://" + u
    return u.rstrip("/")


# ---------- load scan results ----------
res = {}
with open(os.path.join(D, "results.jsonl")) as f:
    for line in f:
        try:
            r = json.loads(line)
        except Exception:
            continue
        res[r["url"]] = r      # later lines win (re-checks)

print(f"scan results loaded: {len(res)}")

VERDICT_TEXT = {"Yes": "Yes", "Likely": "Likely", "No": "No (no public evidence)",
                "Unknown": "Unknown (unreachable/blocked)"}
FILL = {
    "Yes": PatternFill("solid", fgColor="D5EBD8"),
    "Likely": PatternFill("solid", fgColor="E8F0D9"),
    "No": PatternFill("solid", fgColor="F2F2F2"),
    "Unknown": PatternFill("solid", fgColor="FCE9D6"),
    "": PatternFill("solid", fgColor="FFFFFF"),
}
FONT = {
    "Yes": Font(name="Calibri", size=11, bold=True, color="1B5E20"),
    "Likely": Font(name="Calibri", size=11, color="4C6B1F"),
    "No": Font(name="Calibri", size=11, color="616161"),
    "Unknown": Font(name="Calibri", size=11, color="8A5300"),
    "": Font(name="Calibri", size=11, color="9E9E9E"),
}

wb_src = load_workbook(SRC)
ws_src = wb_src[SHEET]
n_cols_src = ws_src.max_column
n_rows = ws_src.max_row

wb = Workbook()
wb.remove(wb.active)
ws = wb.create_sheet(SHEET)

EXTRA_HEAD = ["Stripe evidence", "Other payment tech seen", "Check status", "URL checked"]
counts = {"Yes": 0, "Likely": 0, "No": 0, "Unknown": 0, "": 0}
evidence_tally = {}
others_tally = {}

for r in range(1, n_rows + 1):
    for c in range(1, n_cols_src + 1):
        sc = ws_src.cell(row=r, column=c)
        dc_idx = c if c < INSERT_AT else c + 1
        dc = ws.cell(row=r, column=dc_idx, value=sc.value)
        dc._style = copy(sc._style)
        if sc.hyperlink is not None:
            dc.hyperlink = sc.hyperlink.target

    if r == 1:
        h = ws.cell(row=1, column=INSERT_AT, value="Stripe?")
        h._style = copy(ws_src.cell(row=1, column=1)._style)
        for i, name in enumerate(EXTRA_HEAD):
            hh = ws.cell(row=1, column=n_cols_src + 2 + i, value=name)
            hh._style = copy(ws_src.cell(row=1, column=1)._style)
        continue

    site = ws_src.cell(row=r, column=WEBSITE_COL).value
    key = norm(site if isinstance(site, str) else "")
    rec = res.get(key)
    if not key:
        verdict, ev, others, status, checked_url = "", "", "", "no website in row", ""
    elif rec is None:
        verdict, ev, others, status, checked_url = "Unknown", "", "", "not checked", ""
    else:
        verdict = rec.get("verdict", "Unknown")
        ev = rec.get("evidence", "") or ""
        others = rec.get("others", "") or ""
        status = rec.get("status", "") or ""
        checked_url = rec.get("found_on", "") or ""
        if verdict in ("Yes", "Likely"):
            for part in ev.split(";"):
                p = part.strip()
                if p:
                    evidence_tally[p] = evidence_tally.get(p, 0) + 1
        for o in others.split(","):
            o = o.strip()
            if o:
                others_tally[o] = others_tally.get(o, 0) + 1

    counts[verdict] = counts.get(verdict, 0) + 1
    cell = ws.cell(row=r, column=INSERT_AT, value=VERDICT_TEXT.get(verdict, ""))
    cell.fill = FILL.get(verdict, FILL[""])
    cell.font = FONT.get(verdict, FONT[""])
    cell.alignment = Alignment(horizontal="center", vertical="center")

    for i, val in enumerate((ev, others, status, checked_url)):
        c2 = ws.cell(row=r, column=n_cols_src + 2 + i, value=val)
        c2.font = Font(name="Calibri", size=9, color="595959")

# widths, panes, filter
for c in range(1, n_cols_src + 1):
    letter_src = get_column_letter(c)
    dim = ws_src.column_dimensions.get(letter_src)
    if dim is not None and dim.width:
        ws.column_dimensions[get_column_letter(c if c < INSERT_AT else c + 1)].width = dim.width
ws.column_dimensions[get_column_letter(INSERT_AT)].width = 22
for i, w in enumerate((44, 30, 18, 46)):
    ws.column_dimensions[get_column_letter(n_cols_src + 2 + i)].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(n_cols_src + 1 + len(EXTRA_HEAD) + 1)}1"
ws.cell(row=1, column=INSERT_AT).comment = Comment(
    "Yes = concrete Stripe marker found on a public page (see 'Stripe evidence').\n"
    "No (no public evidence) = nothing found on homepage + pricing/checkout pages + main JS bundles; "
    "Stripe may still be used behind a login.\n"
    "Unknown = site unreachable, DNS-dead, or bot-blocked (403/429).", "Stripe check")

# ---------- copy Summary sheet ----------
if "Summary" in wb_src.sheetnames:
    s_src = wb_src["Summary"]
    s = wb.create_sheet("Summary")
    for r in range(1, s_src.max_row + 1):
        for c in range(1, s_src.max_column + 1):
            sc = s_src.cell(row=r, column=c)
            dc = s.cell(row=r, column=c, value=sc.value)
            dc._style = copy(sc._style)
    for c in range(1, s_src.max_column + 1):
        L = get_column_letter(c)
        d = s_src.column_dimensions.get(L)
        if d is not None and d.width:
            s.column_dimensions[L].width = d.width

# ---------- methodology sheet ----------
m = wb.create_sheet("Stripe Check Method")
m.column_dimensions["A"].width = 116
m.column_dimensions["B"].width = 14
B = Font(name="Calibri", size=12, bold=True)
N = Font(name="Calibri", size=10)
total_sites = counts["Yes"] + counts["Likely"] + counts["No"] + counts["Unknown"]
lines = [
    ("STRIPE DETECTION - method and results", 14, True, None),
    (f"Checked {total_sites:,} rows with a website ({len(res):,} unique domains actually fetched). "
     f"{counts['']:,} rows have no website value.", 10, False, None),
    ("", 10, False, None),
    ("RESULTS", 12, True, None),
    ("Stripe confirmed (Yes) - hard technical artifact found", 10, False, counts["Yes"]),
    ("Likely - Stripe named/referenced but no hard artifact on public pages", 10, False, counts["Likely"]),
    ("No public evidence of Stripe", 10, False, counts["No"]),
    ("Unknown (unreachable, DNS-dead, or bot-blocked)", 10, False, counts["Unknown"]),
    ("Rows with no website listed", 10, False, counts[""]),
    ("", 10, False, None),
    ("HOW EACH SITE WAS CHECKED", 12, True, None),
    ("1. GET the homepage over https (falling back to http), following redirects, browser user-agent.", 10, False, None),
    ("2. Scan the HTML and the response headers (incl. Content-Security-Policy) for concrete Stripe markers.", 10, False, None),
    ("3. If nothing found: follow up to 2 commerce links found on the homepage (pricing / plans / checkout /", 10, False, None),
    ("   subscribe / billing / upgrade / buy / signup), plus /pricing and /plans directly, and scan those.", 10, False, None),
    ("4. If still nothing: download up to 2 first-party JavaScript bundles and scan those (catches SPAs that", 10, False, None),
    ("   load Stripe.js from application code rather than a script tag).", 10, False, None),
    ("", 10, False, None),
    ("WHAT COUNTS AS 'YES' EVIDENCE (no guessing, no keyword-matching on the bare word 'stripe')", 12, True, None),
    ("js.stripe.com  |  checkout.stripe.com  |  buy.stripe.com (payment links)  |  billing.stripe.com (portal)", 10, False, None),
    ("api.stripe.com  |  m.stripe.network (fraud beacon)  |  stripecdn.com  |  q.stripe.com  |  hooks.stripe.com", 10, False, None),
    ("pk_live_ / pk_test_ publishable keys  |  <stripe-pricing-table>  |  <stripe-buy-button>  |  data-stripe attributes", 10, False, None),
    ("@stripe/stripe-js package reference  |  Stripe(pk_...) initialisation  |  Stripe.js API calls  |  WooCommerce Stripe plugin", 10, False, None),
    ("The bare word 'stripe' is deliberately NOT matched, so CSS classes like .striped and sentences such as", 10, False, None),
    ("'we integrate with Zapier, Stripe and Slack' do NOT produce a false positive (verified by unit test).", 10, False, None),
    ("", 10, False, None),
    ("WHAT COUNTS AS 'LIKELY' (weaker, shown separately so you can audit it)", 12, True, None),
    ("'Powered by Stripe' / 'payments processed by Stripe' text on the site's own pricing or checkout page;", 10, False, None),
    ("application code referencing stripeSession / stripeCustomer / data-stripe; API routes named .../stripe.", 10, False, None),
    ("", 10, False, None),
    ("LIMITS YOU SHOULD KNOW (read before acting on a 'No')", 12, True, None),
    ("A 'No' means no Stripe evidence was visible on the public pages checked - NOT proof the company has no", 10, False, None),
    ("Stripe account. Companies whose billing lives behind a login (verified examples: Figma, Vercel, Calendly)", 10, False, None),
    ("do use Stripe but show no public trace. Treat 'Yes' as high-confidence and 'No' as 'not publicly visible'.", 10, False, None),
    ("Mobile-app rows often bill through Apple/Google in-app purchase instead - see 'Other payment tech seen'.", 10, False, None),
    ("Sites behind Cloudflare/bot protection return 403 and are marked Unknown rather than guessed.", 10, False, None),
    ("A site can also load Stripe only after a user clicks 'Subscribe', which a static fetch never triggers.", 10, False, None),
    ("", 10, False, None),
    ("MOST COMMON EVIDENCE MARKERS", 12, True, None),
]
for k, v in sorted(evidence_tally.items(), key=lambda kv: -kv[1])[:12]:
    lines.append((f"   {k}", 10, False, v))
lines.append(("", 10, False, None))
lines.append(("OTHER PAYMENT TECH SEEN ACROSS THE LIST", 12, True, None))
for k, v in sorted(others_tally.items(), key=lambda kv: -kv[1])[:16]:
    lines.append((f"   {k}", 10, False, v))

for i, (txt, size, bold, num) in enumerate(lines, 1):
    c = m.cell(row=i, column=1, value=txt)
    c.font = Font(name="Calibri", size=size, bold=bold)
    if num is not None:
        nc = m.cell(row=i, column=2, value=num)
        nc.font = Font(name="Calibri", size=size, bold=bold)
        nc.number_format = "#,##0"

wb.save(OUT)
print("saved", OUT)
print("counts:", counts)
print("top evidence:", sorted(evidence_tally.items(), key=lambda kv: -kv[1])[:6])
print("top others:", sorted(others_tally.items(), key=lambda kv: -kv[1])[:8])
